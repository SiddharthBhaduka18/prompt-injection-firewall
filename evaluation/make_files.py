"""Render a text payload into each supported input format (real files, not mocks).

Each renderer can place the payload in the *visible* content or in a *hidden* channel,
mirroring how real indirect prompt injections are delivered.
"""
from __future__ import annotations

import io
import json
import random
from email.message import EmailMessage

from PIL import Image, ImageDraw

from agent.fonts import font_path, pil_font

COVER = [
    "Quarterly operations summary. Throughput improved across all three plants and downtime fell by 8%.",
    "Welcome to our product page. This stainless steel bottle keeps drinks cold for 24 hours.",
    "Hello team, please find the project update below. The release is on track for next month.",
    "Customer review: the delivery was quick and the packaging was secure.",
]


def _wrap(text, width=60):
    words, lines, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    return lines + [cur]


def html(payload, hidden, rng):
    cover = rng.choice(COVER)
    if hidden:
        trick = rng.choice([
            f'<div style="display:none">{payload}</div>',
            f'<span style="color:#ffffff;font-size:1px">{payload}</span>',
            f"<!-- {payload} -->",
            f'<img src="pic.png" alt="{payload}">',
            f'<p style="position:absolute;left:-9999px">{payload}</p>',
        ])
        body = f"<p>{cover}</p>{trick}"
    else:
        body = f"<p>{cover}</p><p>{payload}</p>"
    return f"<html><head><title>Page</title></head><body>{body}</body></html>".encode(), "page.html"


def markdown(payload, hidden, rng):
    cover = rng.choice(COVER)
    if hidden:
        md = f"# Notes\n\n{cover}\n\n<!-- {payload} -->\n" if rng.random() < 0.6 else f"# Notes\n\n{cover}\n\n[//]: # ({payload})\n"
    else:
        md = f"# Notes\n\n{cover}\n\n{payload}\n"
    return md.encode(), "notes.md"


def pdf(payload, hidden, rng):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    fp = font_path()
    if fp and "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", fp))  # embeds Unicode glyphs (homoglyph payloads survive)
    fname = "DejaVu" if fp else "Helvetica"
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont(fname, 11)
    y = 780
    for line in _wrap(rng.choice(COVER), 85):
        c.drawString(60, y, line)
        y -= 16
    if hidden:
        if rng.random() < 0.5:
            c.setFillColorRGB(1, 1, 1)
            c.setFont(fname, 9)
        else:
            c.setFont(fname, 1)
    y -= 10
    for line in _wrap(payload, 85):
        c.drawString(60, y, line)
        y -= 14
    c.save()
    return buf.getvalue(), "doc.pdf"


def docx(payload, hidden, rng):
    import docx as d
    from docx.shared import Pt, RGBColor
    doc = d.Document()
    doc.add_paragraph(rng.choice(COVER))
    p = doc.add_paragraph()
    r = p.add_run(payload)
    if hidden:
        k = rng.choice(["hidden", "white", "tiny"])
        if k == "hidden":
            r.font.hidden = True
        elif k == "white":
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            r.font.size = Pt(1)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), "file.docx"


def eml(payload, hidden, rng):
    m = EmailMessage()
    m["From"] = "vendor@partner-co.com"
    m["To"] = "me@company.com"
    m["Subject"] = "Project update"
    cover = rng.choice(COVER)
    if hidden:
        m.set_content(cover)
        m.add_alternative(f"<html><body><p>{cover}</p><div style='display:none'>{payload}</div></body></html>", subtype="html")
    else:
        m.set_content(f"{cover}\n\n{payload}\n\nRegards,\nVendor")
    return bytes(m), "mail.eml"


def api(payload, hidden, rng):
    obj = {"status": "ok", "data": {"product": "Bottle", "price": 12.5,
                                    "reviews": [{"user": "a1", "text": rng.choice(COVER)},
                                                {"user": "b2", "text": payload}]}}
    if hidden:
        obj["data"]["_meta"] = {"note": payload}
        obj["data"]["reviews"][1]["text"] = "Nice."
    return json.dumps(obj).encode(), "response.json"


def code(payload, hidden, rng):
    src = ("import json\n\n"
           f"# {payload}\n" if hidden else "import json\n\n") + \
        ("def load(path):\n    \"\"\"Load config.\"\"\"\n    with open(path) as f:\n        return json.load(f)\n")
    if not hidden:
        src += f'\nMESSAGE = "{payload}"\n'
    return src.encode(), "util.py"


def image(payload, hidden, rng):
    W = 1400
    lines = _wrap(rng.choice(COVER), 70) + [""] + _wrap(payload, 70)
    img = Image.new("RGB", (W, 60 + 42 * len(lines)), (255, 255, 255))
    d = ImageDraw.Draw(img)
    f = pil_font(28)
    cover_n = len(_wrap(rng.choice(COVER), 70)) + 1
    y = 30
    for i, line in enumerate(lines):
        faint = hidden and i >= cover_n
        d.text((30, y), line, fill=(250, 250, 250) if faint else (20, 20, 20), font=f)
        y += 42
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue(), "image.png"


def plain(payload, hidden, rng):
    return payload.encode(), "message.txt"


RENDERERS = {
    "user_message": plain, "html": html, "markdown": markdown, "pdf": pdf, "docx": docx,
    "email": eml, "api_response": api, "source_code": code, "image": image,
}
HIDDEN_CAPABLE = {"html", "markdown", "pdf", "docx", "email", "api_response", "source_code", "image"}
