"""Layer 0 - Multimodal ingestion.

Turns any supported input into a list of `Segment`s. Crucially it separates what a
human sees from *hidden channels* an LLM still reads (CSS-hidden HTML, comments,
white/tiny PDF text, hidden Word runs, metadata, alt text, faint text in images...).
Instructions in a hidden channel are a strong signal of indirect injection.

Supported: user message, web page / HTML, Markdown, PDF, email (.eml), Word (.docx),
API/JSON responses, OCR text, source code, images (PNG/JPG via Tesseract OCR).
"""
from __future__ import annotations

import email
import io
import json
import os
import re
from dataclasses import dataclass
from email import policy

SOURCE_TYPES = ["user_message", "web_page", "html", "markdown", "pdf", "email", "docx",
                "api_response", "ocr_text", "source_code", "image"]


@dataclass
class Segment:
    text: str
    channel: str          # e.g. visible, html-hidden-css, pdf-white-text, docx-comment ...
    hidden: bool = False  # not visible to a human reader

    def __post_init__(self):
        self.text = (self.text or "").strip()


def detect_type(filename: str, data: bytes | None = None) -> str:
    ext = os.path.splitext(filename.lower())[1]
    return {
        ".html": "html", ".htm": "web_page", ".md": "markdown", ".markdown": "markdown",
        ".pdf": "pdf", ".eml": "email", ".docx": "docx", ".json": "api_response",
        ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image", ".bmp": "image",
        ".py": "source_code", ".js": "source_code", ".ts": "source_code", ".java": "source_code",
        ".go": "source_code", ".c": "source_code", ".cpp": "source_code", ".rb": "source_code",
        ".sh": "source_code", ".txt": "ocr_text" if "ocr" in filename.lower() else "user_message",
    }.get(ext, "user_message")


# ----------------------------------------------------------------- HTML / web
HIDDEN_CSS = re.compile(
    r"display\s*:\s*none|visibility\s*:\s*hidden|opacity\s*:\s*0(?:\.0+)?\s*(?:;|$)|font-size\s*:\s*0(?:px|pt|em)?\s*(?:;|$)"
    r"|font-size\s*:\s*[01](?:\.\d+)?px|color\s*:\s*(?:#fff(?:fff)?|white|rgba?\(\s*255\s*,\s*255\s*,\s*255)"
    r"|(?:left|top|text-indent)\s*:\s*-\d{3,}|height\s*:\s*0|width\s*:\s*0|clip\s*:\s*rect\(0", re.I)


def from_html(raw: str) -> list[Segment]:
    from bs4 import BeautifulSoup, Comment
    soup = BeautifulSoup(raw, "html.parser")
    segs: list[Segment] = []
    for c in soup.find_all(string=lambda s: isinstance(s, Comment)):
        segs.append(Segment(str(c), "html-comment", True))
        c.extract()
    for m in soup.find_all("meta"):
        if m.get("content") and len(m.get("content", "")) > 15:
            segs.append(Segment(m["content"], "html-meta", True))
    for s in soup.find_all("script"):
        if s.get("type", "").endswith("ld+json") or (s.string and re.search(r"[a-z]{4,} [a-z]{4,} [a-z]{4,}", s.string or "")):
            segs.append(Segment(s.get_text(), "html-script", True))
        s.decompose()
    for s in soup.find_all("style"):
        s.decompose()
    for el in soup.find_all(True):
        for attr in ("alt", "title", "aria-label", "data-prompt", "data-instructions"):
            v = el.get(attr)
            if v and len(v) > 15:
                segs.append(Segment(v, f"html-attr-{attr}", attr != "alt"))
    hidden_els = [el for el in soup.find_all(True)
                  if el.has_attr("hidden") or el.get("aria-hidden") == "true"
                  or HIDDEN_CSS.search(el.get("style", "") or "")
                  or el.name in ("noscript", "template")]
    for el in hidden_els:
        if el.parent is None:  # already removed with an ancestor
            continue
        txt = el.get_text(" ", strip=True)
        if txt:
            segs.append(Segment(txt, "html-hidden-css", True))
        el.decompose()
    visible = soup.get_text(" ", strip=True)
    return [Segment(visible, "visible")] + segs


# ----------------------------------------------------------------- Markdown
def from_markdown(raw: str) -> list[Segment]:
    segs = []
    for m in re.finditer(r"<!--(.*?)-->", raw, re.S):
        segs.append(Segment(m.group(1), "md-comment", True))
    for m in re.finditer(r"^\[//\]:\s*#\s*[(\"'](.*?)[)\"']\s*$", raw, re.M):
        segs.append(Segment(m.group(1), "md-hidden-reference", True))
    for m in re.finditer(r"!\[([^\]]{10,})\]\(", raw):
        segs.append(Segment(m.group(1), "md-image-alt", True))
    for m in re.finditer(r"\[[^\]]*\]\([^)\s]+\s+\"([^\"]{10,})\"\)", raw):
        segs.append(Segment(m.group(1), "md-link-title", True))
    visible = re.sub(r"<!--.*?-->", "", raw, flags=re.S)
    visible = re.sub(r"^\[//\]:.*$", "", visible, flags=re.M)
    if re.search(r"<[a-z][^>]*>", visible, re.I):
        html_segs = from_html(visible)
        return html_segs + segs
    # keep link/image URLs visible to the scanner (exfil beacons live there)
    return [Segment(visible, "visible")] + segs


# ----------------------------------------------------------------- PDF
def from_pdf(data: bytes) -> list[Segment]:
    import pdfplumber
    from pypdf import PdfReader
    segs: list[Segment] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            vis, hid = [], []
            for ch in page.chars:
                col = ch.get("non_stroking_color")
                white = col is not None and (
                    (isinstance(col, (int, float)) and col >= 0.95)
                    or (isinstance(col, (list, tuple)) and len(col) in (1, 3) and all(c >= 0.95 for c in col))
                    or (isinstance(col, (list, tuple)) and len(col) == 4 and all(c <= 0.05 for c in col)))
                tiny = ch.get("size", 10) < 2.0
                off = ch["x0"] > page.width or ch["top"] > page.height or ch["x1"] < 0
                (hid if (white or tiny or off) else vis).append(ch)
            # re-join soft line wraps so sentences split across lines are scanned whole
            vt = re.sub(r"(?<!\n)\n(?!\n)", " ", _chars_to_text(vis))
            ht = re.sub(r"(?<!\n)\n(?!\n)", " ", _chars_to_text(hid))
            if vt:
                segs.append(Segment(vt, f"visible (page {pno})"))
            if ht:
                segs.append(Segment(ht, f"pdf-invisible-text (page {pno})", True))
            for a in page.annots or []:
                c = (a.get("data") or {}).get("Contents")
                if c:
                    segs.append(Segment(c if isinstance(c, str) else c.decode("latin-1", "ignore"), "pdf-annotation", True))
    try:
        meta = PdfReader(io.BytesIO(data)).metadata or {}
        mt = " | ".join(str(v) for k, v in meta.items() if v and len(str(v)) > 12)
        if mt:
            segs.append(Segment(mt, "pdf-metadata", True))
    except Exception:
        pass
    if not any(s.text for s in segs):  # scanned PDF -> OCR
        segs += _ocr_pdf(data)
    return segs


def _chars_to_text(chars) -> str:
    if not chars:
        return ""
    chars = sorted(chars, key=lambda c: (round(c["top"]), c["x0"]))
    out, last_top, last_x1 = [], None, None
    for c in chars:
        if last_top is not None and abs(c["top"] - last_top) > 3:
            out.append("\n")
        elif last_x1 is not None and c["x0"] - last_x1 > c.get("size", 10) * 0.15:
            out.append(" ")
        out.append(c["text"])
        last_top, last_x1 = c["top"], c["x1"]
    return "".join(out)


def _ocr_pdf(data: bytes) -> list[Segment]:
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(data)
        segs = []
        for i in range(min(len(doc), 5)):
            img = doc[i].render(scale=2).to_pil()
            segs += [Segment(s.text, f"ocr (page {i+1}) " + s.channel, s.hidden) for s in _ocr_image(img)]
        return segs
    except Exception as e:
        return [Segment("", f"ocr-failed: {e}")]


# ----------------------------------------------------------------- DOCX
def from_docx(data: bytes) -> list[Segment]:
    import docx
    from docx.shared import RGBColor
    d = docx.Document(io.BytesIO(data))
    vis, hid = [], []

    def runs_of(paragraphs):
        for p in paragraphs:
            for r in p.runs:
                yield r

    all_pars = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                all_pars += cell.paragraphs
    for p in all_pars:
        pv, ph = [], []
        for r in p.runs:
            f = r.font
            hidden = bool(f.hidden) or (f.size is not None and f.size.pt < 2) or \
                (f.color is not None and f.color.type is not None and f.color.rgb in (RGBColor(0xFF, 0xFF, 0xFF),))
            (ph if hidden else pv).append(r.text)
        if "".join(pv).strip():
            vis.append("".join(pv))
        if "".join(ph).strip():
            hid.append("".join(ph))
    segs = [Segment("\n".join(vis), "visible")]
    if hid:
        segs.append(Segment("\n".join(hid), "docx-hidden-text", True))
    for sec in d.sections:
        for part, name in ((sec.header, "docx-header"), (sec.footer, "docx-footer")):
            txt = "\n".join(p.text for p in part.paragraphs if p.text.strip())
            if txt:
                segs.append(Segment(txt, name))
    # comments part
    try:
        for rel in d.part.rels.values():
            if "comments" in rel.reltype:
                xml = rel.target_part.blob.decode("utf-8", "ignore")
                txt = " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
                if txt.strip():
                    segs.append(Segment(txt, "docx-comment", True))
    except Exception:
        pass
    cp = d.core_properties
    meta = " | ".join(x for x in (cp.title, cp.subject, cp.comments, cp.keywords, cp.author) if x and len(x) > 12)
    if meta:
        segs.append(Segment(meta, "docx-metadata", True))
    return segs


# ----------------------------------------------------------------- Email
def from_email(data: bytes) -> list[Segment]:
    msg = email.message_from_bytes(data, policy=policy.default)
    segs = [Segment(f"From: {msg['from']} | Subject: {msg['subject']}", "email-header")]
    for part in msg.walk():
        ctype = part.get_content_type()
        disp = part.get_content_disposition()
        if part.is_multipart():
            continue
        payload = part.get_payload(decode=True) or b""
        if disp == "attachment":
            fname = part.get_filename() or "attachment"
            for s in extract(payload, detect_type(fname), fname):
                segs.append(Segment(s.text, f"email-attachment:{fname} / {s.channel}", s.hidden))
        elif ctype == "text/plain":
            segs.append(Segment(payload.decode(part.get_content_charset() or "utf-8", "ignore"), "visible"))
        elif ctype == "text/html":
            for s in from_html(payload.decode(part.get_content_charset() or "utf-8", "ignore")):
                segs.append(Segment(s.text, "email-html " + s.channel if s.channel != "visible" else "email-html", s.hidden))
    return segs


# ----------------------------------------------------------------- JSON / API
def from_json(raw: str) -> list[Segment]:
    try:
        obj = json.loads(raw)
    except json.JSONDecodeError:
        return [Segment(raw, "visible")]
    segs = []

    def walk(o, path="$"):
        if isinstance(o, dict):
            for k, v in o.items():
                walk(v, f"{path}.{k}")
        elif isinstance(o, list):
            for i, v in enumerate(o):
                walk(v, f"{path}[{i}]")
        elif isinstance(o, str) and len(o) > 3:
            segs.append(Segment(o, f"json {path}"))
    walk(obj)
    return segs or [Segment(raw, "visible")]


# ----------------------------------------------------------------- Source code
def from_code(raw: str) -> list[Segment]:
    comments = re.findall(r"(?m)(?:#|//)\s?(.+)$", raw)
    comments += re.findall(r"/\*(.*?)\*/", raw, re.S)
    docstrings = re.findall(r"(?:\"\"\"|''')(.*?)(?:\"\"\"|''')", raw, re.S)
    strings = [s for s in re.findall(r"[\"']([^\"'\n]{25,})[\"']", raw)]
    segs = [Segment(raw, "code")]
    if comments:
        segs.append(Segment("\n".join(comments), "code-comments", True))
    if docstrings:
        segs.append(Segment("\n".join(docstrings), "code-docstrings", True))
    if strings:
        segs.append(Segment("\n".join(strings), "code-string-literals", True))
    return segs


# ----------------------------------------------------------------- Images
def _ocr_image(img) -> list[Segment]:
    import pytesseract
    from PIL import ImageOps, ImageFilter
    segs = []
    g = img.convert("L")
    normal = pytesseract.image_to_string(g)
    segs.append(Segment(normal, "ocr-visible"))
    # Contrast-stretch to reveal faint / near-background text that humans can't see
    # but vision-language models can read.
    stretched = ImageOps.autocontrast(g, cutoff=1)
    hist = stretched.histogram()
    bg = max(range(256), key=lambda i: hist[i])
    lut = [255 if abs(i - bg) <= 1 else 0 for i in range(256)]
    revealed = stretched.point(lut).filter(ImageFilter.MedianFilter(3))
    faint = pytesseract.image_to_string(revealed)
    # second pass on the raw image: any pixel that differs from the background at all becomes ink
    hist_raw = g.histogram()
    bg_raw = max(range(256), key=lambda i: hist_raw[i])
    exact = g.point([255 if i == bg_raw else 0 for i in range(256)]).filter(ImageFilter.MedianFilter(3))
    faint2 = pytesseract.image_to_string(exact)
    if len(set(re.findall(r"[a-z]{4,}", faint2.lower()))) > len(set(re.findall(r"[a-z]{4,}", faint.lower()))):
        faint = faint2
    extra = _new_words(faint, normal)
    if extra:
        segs.append(Segment(faint, "ocr-low-contrast-text", True))
    return segs


def _new_words(a: str, b: str) -> bool:
    wa = set(re.findall(r"[a-z]{4,}", a.lower()))
    wb = set(re.findall(r"[a-z]{4,}", b.lower()))
    return len(wa - wb) >= 3


def from_image(data: bytes) -> list[Segment]:
    from PIL import Image
    img = Image.open(io.BytesIO(data))
    segs = []
    meta = " | ".join(f"{k}: {v}" for k, v in (img.info or {}).items()
                      if isinstance(v, str) and len(v) > 12)
    try:
        exif = img.getexif()
        meta += " | ".join(str(v) for v in exif.values() if isinstance(v, str) and len(v) > 12)
    except Exception:
        pass
    if img.mode not in ("L", "RGB"):
        img = img.convert("RGB")
    segs += _ocr_image(img)
    if meta.strip(" |"):
        segs.append(Segment(meta, "image-metadata", True))
    return segs


# ----------------------------------------------------------------- dispatcher
def extract(data: bytes | str, source_type: str, filename: str = "") -> list[Segment]:
    as_text = data if isinstance(data, str) else None

    def text():
        return as_text if as_text is not None else data.decode("utf-8", "ignore")

    try:
        if source_type in ("html", "web_page"):
            segs = from_html(text())
        elif source_type == "markdown":
            segs = from_markdown(text())
        elif source_type == "pdf":
            segs = from_pdf(data if isinstance(data, bytes) else data.encode())
        elif source_type == "docx":
            segs = from_docx(data)
        elif source_type == "email":
            segs = from_email(data if isinstance(data, bytes) else data.encode())
        elif source_type == "api_response":
            segs = from_json(text())
        elif source_type == "source_code":
            segs = from_code(text())
        elif source_type == "image":
            segs = from_image(data)
        else:  # user_message, ocr_text, plain text
            segs = [Segment(text(), "visible")]
    except Exception as e:  # fault tolerance: never crash the pipeline on a bad file
        fallback = text() if isinstance(data, str) else data[:20000].decode("utf-8", "ignore")
        segs = [Segment(fallback, f"raw-fallback ({type(e).__name__})")]
    return [s for s in segs if s.text]
